# -*- coding: utf-8 -*-
r"""
================================================================
 comun/excel.py  -  RESULTADOS EN EXCEL
================================================================
 Escribe libros .xlsx que una persona puede abrir sin saber nada del
 repositorio: una hoja LEEME que explica el resto, encabezados con su
 unidad, la fila de titulos fija, filtros, anchos y formatos numericos.

   escribir_libro_anexo(ed, ruta, argv=None)            el del edificio
   escribir_libro_reanalisis(modelo, respuesta, ruta)   el de /analizar

 Las firmas las fija semana05/CONTRATO.md (seccion c): el servidor
 (comun/servidor_opensees.py) y el CLI (semana05/exportar_excel.py) las
 llaman tal cual.

 ----------------------------------------------------------------
 EN EXCEL NO SE CALCULA NADA
 ----------------------------------------------------------------
 Ninguna celda tiene formula. Todo numero sale de Python: el anexo de
 semana04/exportar_unity.construir_anexo (los mismos numeros que
 muestra Unity), la respuesta del servidor y calcular.equilibrio. Un
 texto que empieza con '=' se escribe como texto, no como formula.
 Si alguien cambia una celda no cambia nada: se regenera el libro.

 ----------------------------------------------------------------
 DOS PARTES
 ----------------------------------------------------------------
 1. GENERICA (Columna, Hoja, escribir, leer, estructura): no sabe de
    edificios. Sirve para cualquier tabla de resultados -- la idea es
    reutilizarla en la tesis --. Garantiza:
      - escritura ATOMICA: se escribe un temporal y se reemplaza con
        os.replace; un corte a medias no deja un .xlsx roto;
      - un PermissionError con mensaje claro si el archivo esta abierto
        en Excel (Windows no deja reemplazar un archivo abierto);
      - un libro DETERMINISTA: mismos datos, mismos bytes. openpyxl
        graba la hora actual en docProps/core.xml y en cada entrada del
        ZIP; data/excel/ va a git, y sin esto cada regeneracion seria un
        diff binario aunque nada hubiera cambiado;
      - ni NaN ni inf: xlsx no los representa y un centinela escrito
        como numero (9999) se pinta, se ordena y se promedia como dato.
 2. DEL LABORATORIO: arma las hojas desde el anexo y desde la respuesta
    del servidor.

 ----------------------------------------------------------------
 LAS REACCIONES NO SE SUMAN ENTERAS
 ----------------------------------------------------------------
 nodeReaction en un nodo de diafragma trae la fuerza de la restriccion,
 que es interna. Sumar la columna entera dobla el corte basal (LT2 EX:
 -7266.13 kN contra 3633.06 aplicados). Cada fila dice si cuenta en
 Fx/Fy y en Fz con la regla de calcular.equilibrio, y el total correcto
 esta en la hoja Resumen, calculado por calcular.equilibrio.
 semana05/test_excel.py comprueba que la suma FILTRADA de la hoja es
 ese total: si alguien cambia la regla alla, el test lo dice.
================================================================
"""
from __future__ import annotations

import datetime
import hashlib
import importlib.util
import io
import json
import math
import os
import re
import sys
import threading
import zipfile
import xml.etree.ElementTree as ET

_AQUI = os.path.dirname(os.path.abspath(__file__))
if _AQUI not in sys.path:
    sys.path.insert(0, _AQUI)

import rutas                                 # noqa: E402

# openpyxl se importa dentro de las funciones que escriben o leen: asi
# 'import excel' no falla en una maquina sin la dependencia, y el
# servidor puede devolver el motivo en 'excel_error' sin caerse.


# ============================================================
# 1. PARTE GENERICA
# ============================================================
# La fecha que va en las propiedades del libro y en cada entrada del
# ZIP. Fija a proposito (ver arriba, DETERMINISTA): cuando se genero lo
# dice git, no el archivo.
FECHA_FIJA = datetime.datetime(2026, 1, 1, 0, 0, 0)
FECHA_ZIP = (2026, 1, 1, 0, 0, 0)
AUTOR = 'Grupo 7 - Metodos Computacionales UAndes (escrito por Python)'

# Formatos numericos. Solo cambian lo que se VE: la celda guarda el
# numero entero con todos sus decimales. En Windows en castellano el
# punto decimal se muestra como coma.
F_ENTERO = '0'
F_M = '0.000'
F_MM = '0.000'
F_RAD = '0.000E+00'
F_KN = '0.00'
F_KNM = '0.00'
F_DC = '0.000'
F_LAMBDA = '0.00'
F_AREA = '0.0000'
F_INERCIA = '0.000E+00'
F_KPA = '#,##0'
F_MPA = '0.0'
F_CM2 = '0.00'
F_PCT = '0.00'
F_ERROR = '0.0E+00'

# Colores (los de "relleno rojo claro con texto rojo oscuro" de Excel).
ROJO = ('FFC7CE', '9C0006')
AMBAR = ('FFEB9C', '9C5700')
_AZUL_ENCABEZADO = '1F4E78'
_FONDO_TEXTO = 'F2F2F2'

# Lo que Excel no admite en el nombre de una hoja.
_PROHIBIDOS_HOJA = set('[]:*?/\\')
_MAX_NOMBRE_HOJA = 31
_MAX_ANCHO = 60


class Columna(object):
    """
    Una columna: titulo, unidad (va al encabezado entre corchetes),
    formato numerico de Excel, ancho en caracteres (None = segun el
    contenido) y una nota que la hoja LEEME puede listar.
    """

    def __init__(self, titulo, unidad='', formato=None, ancho=None, nota=''):
        self.titulo = titulo
        self.unidad = unidad
        self.formato = formato
        self.ancho = ancho
        self.nota = nota

    @property
    def encabezado(self):
        return '%s [%s]' % (self.titulo, self.unidad) if self.unidad else self.titulo


class Enlace(object):
    """Un texto que al hacer clic lleva a otra hoja del mismo libro."""

    def __init__(self, texto, hoja):
        self.texto = texto
        self.hoja = hoja


class ReglaColor(object):
    """
    Formato condicional sobre una columna: 'mayor' (valor > a) o 'entre'
    (a <= valor <= b). Excel lo evalua al mostrar; no cambia la celda.
    """

    def __init__(self, columna, operador, a, b=None, color=ROJO):
        if operador not in ('mayor', 'entre'):
            raise ValueError('ReglaColor: operador %r (use mayor o entre)' % operador)
        self.columna = columna
        self.operador = operador
        self.a = a
        self.b = b
        self.color = color


class Hoja(object):
    """
    Una hoja: nombre, columnas y filas (listas del mismo largo que las
    columnas). 'texto' = hoja de lectura (LEEME, Supuestos): ajusta el
    texto dentro de la celda y alinea arriba.
    """

    def __init__(self, nombre, columnas, filas, descripcion='', congelar_columnas=1,
                 filtro=True, reglas=(), texto=False, color_pestana=None):
        self.nombre = nombre
        self.columnas = list(columnas)
        self.filas = [list(f) for f in filas]
        self.descripcion = descripcion
        self.congelar_columnas = congelar_columnas
        self.filtro = filtro
        self.reglas = list(reglas)
        self.texto = texto
        self.color_pestana = color_pestana
        for i, f in enumerate(self.filas):
            if len(f) != len(self.columnas):
                raise ValueError('hoja %r, fila %d: %d valores para %d columnas'
                                 % (nombre, i + 2, len(f), len(self.columnas)))

    def indice(self, titulo):
        """Posicion (0..n-1) de la columna con ese titulo."""
        for i, c in enumerate(self.columnas):
            if c.titulo == titulo:
                return i
        raise KeyError('hoja %r no tiene la columna %r' % (self.nombre, titulo))


def _validar_nombre_hoja(nombre):
    if not nombre or len(nombre) > _MAX_NOMBRE_HOJA:
        raise ValueError('nombre de hoja %r: entre 1 y %d caracteres'
                         % (nombre, _MAX_NOMBRE_HOJA))
    malos = _PROHIBIDOS_HOJA.intersection(nombre)
    if malos:
        raise ValueError('nombre de hoja %r: Excel no admite %s'
                         % (nombre, ' '.join(sorted(malos))))


def _valor(v, donde):
    """
    Lo que se puede escribir en una celda. Rechaza NaN e inf con el
    lugar exacto: preferible a un libro que Excel abre "reparando".
    """
    if v is None or isinstance(v, (str, Enlace)):
        return v
    if isinstance(v, bool):
        return v
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        if not math.isfinite(v):
            raise ValueError('%s: %r no se puede escribir en Excel (use una celda '
                             'vacia y un texto que diga por que)' % (donde, v))
        # -0.0 se veria "-0,00" en Excel: es el mismo cero.
        return v + 0.0
    raise TypeError('%s: tipo %s no admitido (%r)' % (donde, type(v).__name__, v))


def _ancho_auto(col, filas, j):
    """Ancho en caracteres: el encabezado o el texto mas largo, acotado."""
    if col.ancho:
        return col.ancho
    largo = len(col.encabezado) + 3        # + la flecha del filtro
    for f in filas[:500]:
        v = f[j]
        if isinstance(v, Enlace):
            v = v.texto
        if isinstance(v, str):
            largo = max(largo, len(v) + 1)
        elif isinstance(v, float):
            largo = max(largo, 11)
        elif isinstance(v, int) and not isinstance(v, bool):
            largo = max(largo, len(str(v)) + 1)
    return min(largo, _MAX_ANCHO)


def _alto_de_fila(fila, anchos):
    """
    Alto de una fila de texto ajustado: Excel no lo recalcula al abrir
    un libro escrito por otro programa, y sin esto un parrafo quedaria
    cortado en una sola linea.
    """
    lineas = 1
    for v, ancho in zip(fila, anchos):
        if isinstance(v, Enlace):
            v = v.texto
        if not isinstance(v, str) or not v:
            continue
        n = 0
        for parrafo in v.split('\n'):
            n += max(1, int(math.ceil(len(parrafo) / max(1.0, ancho - 1.0))))
        lineas = max(lineas, n)
    return 15.0 * lineas


def mensaje_archivo_abierto(ruta):
    return ('No se pudo escribir %s: el archivo esta abierto en otro programa '
            '(probablemente Excel). Cierralo y vuelve a correr el comando.'
            % os.path.abspath(ruta))


def comprobar_que_se_puede_escribir(ruta):
    r"""
    Falla ANTES de calcular nada si el destino esta abierto en Excel.
    Excel abre el .xlsx negando escritura a los demas y deja al lado un
    '~$<nombre>'. Abrir el archivo para escritura sin truncarlo es la
    prueba directa; el '~$' solo mejora el mensaje.
    """
    ruta = os.path.abspath(ruta)
    if not os.path.exists(ruta):
        return
    try:
        with open(ruta, 'r+b'):
            pass
    except PermissionError:
        raise PermissionError(mensaje_archivo_abierto(ruta))


def _zip_determinista(datos, propiedades):
    """
    Reescribe el ZIP de openpyxl con fecha fija en cada entrada y con
    docProps/core.xml regenerado (save_workbook le pone la hora actual).
    """
    from openpyxl.xml.functions import tostring
    propiedades.created = FECHA_FIJA
    propiedades.modified = FECHA_FIJA
    entrada = zipfile.ZipFile(io.BytesIO(datos))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
        for info in entrada.infolist():
            contenido = entrada.read(info.filename)
            if info.filename == 'docProps/core.xml':
                contenido = tostring(propiedades.to_tree())
            zi = zipfile.ZipInfo(info.filename, date_time=FECHA_ZIP)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            z.writestr(zi, contenido)
    return buf.getvalue()


def _escribir_bytes_atomico(ruta, datos):
    """
    Temporal en la MISMA carpeta + os.replace: el reemplazo es atomico
    en el mismo disco, asi que quien abra el archivo ve el viejo o el
    nuevo, nunca uno a medias.
    """
    ruta = os.path.abspath(ruta)
    rutas.asegurar(ruta)
    carpeta, base = os.path.split(ruta)
    tmp = os.path.join(carpeta, '.%s.%d.%d.tmp' % (base, os.getpid(), threading.get_ident()))
    try:
        with open(tmp, 'wb') as f:
            f.write(datos)
            f.flush()
            os.fsync(f.fileno())
        try:
            os.replace(tmp, ruta)
        except PermissionError:
            raise PermissionError(mensaje_archivo_abierto(ruta))
    finally:
        if os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass
    return ruta


def escribir(ruta, hojas, titulo='', asunto='', descripcion=''):
    r"""
    Escribe las hojas en un .xlsx y devuelve la ruta ABSOLUTA.

    Crea la carpeta si falta. PermissionError con mensaje claro si el
    archivo esta abierto. ValueError si una hoja trae NaN/inf o un
    nombre invalido. La primera hoja queda activa.

    Cada numero se guarda con 16 cifras significativas (openpyxl escribe
    "%.16g"): 16.435000000000002 se relee 16.435. Excel muestra 15.
    """
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.hyperlink import Hyperlink
    from openpyxl.writer.excel import save_workbook

    comprobar_que_se_puede_escribir(ruta)
    nombres = [h.nombre for h in hojas]
    if len(set(n.lower() for n in nombres)) != len(nombres):
        raise ValueError('hojas con nombre repetido: %s' % ', '.join(nombres))
    for n in nombres:
        _validar_nombre_hoja(n)

    wb = Workbook()
    wb.remove(wb.active)

    fuente_encabezado = Font(bold=True, color='FFFFFF')
    relleno_encabezado = PatternFill('solid', fgColor=_AZUL_ENCABEZADO)
    alineado_encabezado = Alignment(wrap_text=True, vertical='center')
    borde_encabezado = Border(bottom=Side(style='thin', color='000000'))
    alineado_texto = Alignment(wrap_text=True, vertical='top')
    fuente_enlace = Font(color='0563C1', underline='single')

    for hoja in hojas:
        ws = wb.create_sheet(hoja.nombre)
        if hoja.color_pestana:
            ws.sheet_properties.tabColor = hoja.color_pestana
        ncol = len(hoja.columnas)

        for j, col in enumerate(hoja.columnas, start=1):
            c = ws.cell(row=1, column=j, value=col.encabezado)
            c.font = fuente_encabezado
            c.fill = relleno_encabezado
            c.alignment = alineado_encabezado
            c.border = borde_encabezado

        anchos = [_ancho_auto(col, hoja.filas, j) for j, col in enumerate(hoja.columnas)]
        formatos = [col.formato for col in hoja.columnas]

        for i, fila in enumerate(hoja.filas, start=2):
            for j in range(ncol):
                donde = 'hoja %r, fila %d, columna %r' % (hoja.nombre, i,
                                                          hoja.columnas[j].encabezado)
                v = _valor(fila[j], donde)
                if v is None:
                    continue
                c = ws.cell(row=i, column=j + 1)
                if isinstance(v, Enlace):
                    c.value = ILLEGAL_CHARACTERS_RE.sub('', v.texto)
                    c.hyperlink = Hyperlink(ref=c.coordinate,
                                            location="'%s'!A1" % v.hoja,
                                            display=v.texto)
                    c.font = fuente_enlace
                elif isinstance(v, str):
                    c.value = ILLEGAL_CHARACTERS_RE.sub('', v)
                    # openpyxl toma '=...' como formula. Aca no hay formulas.
                    if c.data_type == 'f':
                        c.data_type = 's'
                else:
                    c.value = v
                    if formatos[j] and not isinstance(v, bool):
                        c.number_format = formatos[j]
                if hoja.texto:
                    c.alignment = alineado_texto
            if hoja.texto:
                ws.row_dimensions[i].height = _alto_de_fila(fila, anchos)

        for j, ancho in enumerate(anchos, start=1):
            ws.column_dimensions[get_column_letter(j)].width = ancho
        if any(len(col.encabezado) + 3 > ancho for col, ancho in zip(hoja.columnas, anchos)):
            ws.row_dimensions[1].height = 30.0

        ultima = get_column_letter(ncol)
        fin = max(2, len(hoja.filas) + 1)
        if hoja.congelar_columnas is not None:
            # Con la coordenada en texto: ws.cell() crearia una celda vacia
            # fuera de la tabla y la hoja releida traeria una columna de mas.
            ws.freeze_panes = '%s2' % get_column_letter(hoja.congelar_columnas + 1)
        if hoja.filtro and ncol:
            ws.auto_filter.ref = 'A1:%s%d' % (ultima, fin)

        for regla in hoja.reglas:
            letra = get_column_letter(hoja.indice(regla.columna) + 1)
            rango = '%s2:%s%d' % (letra, letra, fin)
            fill = PatternFill('solid', fgColor=regla.color[0], bgColor=regla.color[0])
            font = Font(color=regla.color[1])
            if regla.operador == 'mayor':
                r = CellIsRule(operator='greaterThan', formula=[repr(float(regla.a))],
                               fill=fill, font=font)
            else:
                r = CellIsRule(operator='between',
                               formula=[repr(float(regla.a)), repr(float(regla.b))],
                               fill=fill, font=font)
            ws.conditional_formatting.add(rango, r)

    wb.active = 0
    props = wb.properties
    props.creator = AUTOR
    props.lastModifiedBy = AUTOR
    props.title = titulo
    props.subject = asunto
    props.description = descripcion

    buf = io.BytesIO()
    save_workbook(wb, buf)
    return _escribir_bytes_atomico(ruta, _zip_determinista(buf.getvalue(), props))


def leer(ruta):
    r"""
    {hoja: [filas]} con los VALORES del libro, fila 1 (encabezados)
    incluida. Para los tests: relee lo que se escribio.
    """
    from openpyxl import load_workbook
    wb = load_workbook(ruta, read_only=True, data_only=False)
    try:
        salida = {}
        for ws in wb.worksheets:
            salida[ws.title] = [list(f) for f in ws.iter_rows(values_only=True)]
        return salida
    finally:
        wb.close()


_NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
       'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
       'rel': 'http://schemas.openxmlformats.org/package/2006/relationships'}


def estructura(ruta):
    r"""
    Lo que no son valores, leido del XML del libro sin openpyxl (asi el
    test no depende del mismo lector que el escritor):
      {'hojas': [nombres en orden], 'por_hoja': {nombre: {'panel',
       'filtro', 'reglas' [(rango, operador, [formulas])], 'formulas'
       (cantidad de <f>), 'anchos' (cantidad de <col>), 'enlaces'}},
       'propiedades': texto de docProps/core.xml}
    """
    with zipfile.ZipFile(ruta) as z:
        libro = ET.fromstring(z.read('xl/workbook.xml'))
        rels = ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
        destino = {r.get('Id'): r.get('Target') for r in rels.findall('rel:Relationship', _NS)}
        hojas, por_hoja = [], {}
        for s in libro.find('m:sheets', _NS).findall('m:sheet', _NS):
            nombre = s.get('name')
            hojas.append(nombre)
            objetivo = destino[s.get('{%s}id' % _NS['r'])].lstrip('/')
            if not objetivo.startswith('xl/'):
                objetivo = 'xl/' + objetivo
            x = ET.fromstring(z.read(objetivo))
            panel = x.find('m:sheetViews/m:sheetView/m:pane', _NS)
            filtro = x.find('m:autoFilter', _NS)
            reglas = []
            for cf in x.findall('m:conditionalFormatting', _NS):
                for r in cf.findall('m:cfRule', _NS):
                    reglas.append((cf.get('sqref'), r.get('operator'),
                                   [f.text for f in r.findall('m:formula', _NS)]))
            por_hoja[nombre] = {
                'panel': None if panel is None else dict(panel.attrib),
                'filtro': None if filtro is None else filtro.get('ref'),
                'reglas': reglas,
                'formulas': len(x.findall('.//m:f', _NS)),
                'anchos': len(x.findall('m:cols/m:col', _NS)),
                'enlaces': len(x.findall('m:hyperlinks/m:hyperlink', _NS)),
            }
        core = z.read('docProps/core.xml').decode('utf-8')
    return {'hojas': hojas, 'por_hoja': por_hoja, 'propiedades': core}


# ============================================================
# 2. DEL LABORATORIO
# ============================================================
EDIFICIOS = ('lt2', 'ingenieria', 'conjunto')
CASOS_BASE = ('G', 'Q', 'EX', 'EY')
SI, NO = 'sí', 'no'

# El centinela de semana04/exportar_unity.py (bloque_caso): u = 9999
# cuando Mn = 0 porque P cae fuera de la curva. Es el mismo valor que
# PanelUI.U_FUERA_DE_CURVA en Unity. En el libro va como celda vacia.
U_FUERA_DE_CURVA = 9999.0
ESTADO_PASA = 'pasa'
ESTADO_NO_PASA = 'NO PASA'
ESTADO_FUERA = 'NO PASA: P fuera de la curva'

# Umbrales de COLOR del libro (el coordinador): rojo sobre 1, ambar
# entre 0.9 y 1. Solo pintan; 'pasa' lo decide el anexo (u <= 1).
DC_ROJO = 1.0
DC_AMBAR = 0.9

# El servidor redondea (servidor_opensees.extraer_resultados). Una
# combinacion se redondea igual al escribirla: mas decimales serian
# precision inventada.
DECIMALES_FUERZA = 4
DECIMALES_MM = 5          # 8 decimales en m = 5 en mm

COMANDO_CLI = r'.venv\Scripts\python.exe semana05\exportar_excel.py'
PARAMETROS_JSON = os.path.join(rutas.RAIZ, 'semana03', 'parametros.json')


def _exportador_semana04():
    r"""
    semana04/exportar_unity.py cargado POR RUTA, con el mismo nombre de
    modulo que semana04/trazabilidad.exportador() (si ya lo cargo, se
    reutiliza). 'import exportar_unity' traeria sin error el de la
    Semana 3, que tiene el mismo nombre y no tiene construir_anexo.
    """
    nombre = 'exportar_unity_semana04'
    if nombre in sys.modules:
        return sys.modules[nombre]
    ruta = os.path.join(rutas.RAIZ, 'semana04', 'exportar_unity.py')
    spec = importlib.util.spec_from_file_location(nombre, ruta)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[nombre] = mod
    try:
        spec.loader.exec_module(mod)
    except BaseException:
        del sys.modules[nombre]
        raise
    return mod


def _calcular():
    """comun/calcular.py. Import tardio: calcular importa el servidor, y el
    servidor importa este archivo (ciclo)."""
    import calcular
    return calcular


def _mm(v):
    return round(float(v) * 1000.0, DECIMALES_MM)


def _norma_mm(d):
    return round(1000.0 * math.sqrt(float(d['ux']) ** 2 + float(d['uy']) ** 2
                                    + float(d['uz']) ** 2), DECIMALES_MM)


def _texto_vector(v):
    return '(%s)' % ', '.join('%g' % float(x) for x in v)


def _huella(ruta):
    r"""
    Los primeros 16 del sha256 de un archivo de texto: dice si la fuente
    cambio. Con los fines de linea pasados a LF antes de sumar: git
    (core.autocrlf) entrega data/modelo/*.json con CRLF en Windows y con
    LF en otra maquina, y json.dump escribe CRLF o LF segun el sistema.
    Sin esto el mismo modelo daba otra huella, la hoja LEEME cambiaba,
    test_excel.py [5] declaraba "desactualizado" un libro al dia y
    regenerarlo dejaba un diff binario sin que cambiara ningun numero.
    """
    if not os.path.isfile(ruta):
        return '(no existe)'
    with open(ruta, 'rb') as f:
        contenido = f.read()
    return hashlib.sha256(contenido.replace(b'\r\n', b'\n')).hexdigest()[:16]


def _relativa(ruta):
    try:
        return os.path.relpath(ruta, rutas.RAIZ).replace('\\', '/')
    except ValueError:
        return ruta


# ------------------------------------------------------------
# Nodos, diafragmas y la regla de las reacciones
# ------------------------------------------------------------
def _diafragmas(modelo):
    """(maestros, maestro_de {nodo: maestro}, en_diafragma, atadas).
    maestro_de con el primer diafragma que nombra al nodo, como
    lab_semana03.indice_de_diafragma."""
    maestros, maestro_de, atadas = set(), {}, set()
    for d in modelo.get('diafragmas', []) or []:
        m = int(d['nodo_maestro'])
        maestros.add(m)
        maestro_de.setdefault(m, m)
        for n in d.get('nodos', []) or []:
            maestro_de.setdefault(int(n), m)
        p = int(d.get('perpendicular', 3))
        atadas |= {0: {1, 2}, 1: {1, 2}, 2: {0, 2}, 3: {0, 1}}.get(p, {0, 1})
    return maestros, maestro_de, set(maestro_de), atadas


def _restricciones_declaradas(nd):
    """Las del dato, como las lee el servidor: la lista, o 'fijo' sin
    lista = empotrado. None si el nodo no tiene ninguna."""
    r = nd.get('restricciones') or None
    if r is None and nd.get('fijo', False):
        r = [1, 1, 1, 1, 1, 1]
    if r is None or not any(int(v) for v in r):
        return None
    return [int(v) for v in r]


def _tipo_de_nodo(nd, maestros):
    nid = int(nd['id'])
    if nid in maestros:
        return 'maestro de diafragma'
    if _restricciones_declaradas(nd):
        return 'apoyo'
    if nd.get('auxiliar', False):
        return 'auxiliar'
    return 'libre'


def _hoja_nodos(modelo):
    maestros, maestro_de, _en, _at = _diafragmas(modelo)
    columnas = [
        Columna('Nodo', formato=F_ENTERO),
        Columna('x', 'm', F_M), Columna('y', 'm', F_M), Columna('z', 'm', F_M),
        Columna('Tipo', nota='apoyo (tiene restricciones), maestro de diafragma, '
                             'auxiliar (lo agrega el modelo, p. ej. la cara de un muro) '
                             'o libre'),
        Columna('Restricciones ux uy uz rx ry rz',
                nota='1 = impedido, 0 = libre, tal como vienen en el modelo. Vacio = '
                     'ninguna. El nodo maestro no trae: el solver le fija uz, rx y ry'),
        Columna('Maestro del diafragma', formato=F_ENTERO,
                nota='el nodo maestro del piso rigido al que pertenece; vacio = ninguno'),
    ]
    filas = []
    for nd in sorted(modelo['nodos'], key=lambda n: int(n['id'])):
        nid = int(nd['id'])
        r = _restricciones_declaradas(nd)
        filas.append([nid, float(nd['x']), float(nd['y']), float(nd['z']),
                      _tipo_de_nodo(nd, maestros),
                      ' '.join(str(v) for v in r) if r else None,
                      maestro_de.get(nid)])
    return Hoja('Nodos', columnas, filas,
                descripcion='Un nodo por fila: coordenadas, tipo, restricciones y piso '
                            'rigido al que pertenece.')


def _columnas_desplazamientos():
    return [
        Columna('Caso'), Columna('Nodo', formato=F_ENTERO), Columna('z', 'm', F_M),
        Columna('ux', 'mm', F_MM), Columna('uy', 'mm', F_MM), Columna('uz', 'mm', F_MM),
        Columna('|u|', 'mm', F_MM, nota='norma de la traslacion: raiz(ux² + uy² + uz²)'),
        Columna('rx', 'rad', F_RAD), Columna('ry', 'rad', F_RAD), Columna('rz', 'rad', F_RAD),
    ]


def _filas_desplazamientos(caso, desplazamientos, cota):
    filas = []
    for d in sorted(desplazamientos, key=lambda d: int(d['id'])):
        nid = int(d['id'])
        filas.append([caso, nid, cota.get(nid),
                      _mm(d['ux']), _mm(d['uy']), _mm(d['uz']), _norma_mm(d),
                      float(d['rx']), float(d['ry']), float(d['rz'])])
    return filas


def _columnas_esfuerzos():
    return [
        Columna('Caso'), Columna('Elemento', formato=F_ENTERO), Columna('Tipo'),
        Columna('Extremo', nota='i = nodo n1 (x = 0), j = nodo n2 (x = L)'),
        Columna('x', 'm', F_M),
        Columna('N', 'kN', F_KN, nota='axial interno, TRACCION positiva'),
        Columna('Vy', 'kN', F_KN), Columna('Vz', 'kN', F_KN),
        Columna('T', 'kN·m', F_KNM), Columna('My', 'kN·m', F_KNM), Columna('Mz', 'kN·m', F_KNM),
    ]


def _columnas_reacciones(xy_iguales):
    cols = [
        Columna('Caso'), Columna('Nodo', formato=F_ENTERO),
        Columna('Fx', 'kN', F_KN), Columna('Fy', 'kN', F_KN), Columna('Fz', 'kN', F_KN),
        Columna('Mx', 'kN·m', F_KNM), Columna('My', 'kN·m', F_KNM), Columna('Mz', 'kN·m', F_KNM),
        Columna('Nodo en diafragma', nota='maestro, sí o no'),
    ]
    nota_xy = ('sí = esta fila entra en el total horizontal. Un nodo de diafragma trae '
               'la fuerza INTERNA del piso rigido: no cuenta')
    if xy_iguales:
        cols.append(Columna('Cuenta en Fx/Fy', nota=nota_xy))
    else:
        cols += [Columna('Cuenta en Fx', nota=nota_xy), Columna('Cuenta en Fy', nota=nota_xy)]
    cols.append(Columna('Cuenta en Fz', nota='sí = entra en el total vertical: todo apoyo '
                                             'salvo el nodo maestro'))
    return cols


def _filas_reacciones(caso, reacciones, modelo):
    r"""
    La regla de calcular.equilibrio, fila por fila: en una traslacion
    que ata un diafragma solo cuentan los nodos fuera de todo diafragma;
    el maestro no cuenta en ninguna.
    """
    maestros, _md, en_diafragma, atadas = _diafragmas(modelo)
    xy_iguales = (0 in atadas) == (1 in atadas)

    def cuenta(nid, i):
        return SI if (nid not in maestros
                      and not (i in atadas and nid in en_diafragma)) else NO

    filas = []
    for r in sorted(reacciones, key=lambda r: int(r['id'])):
        nid = int(r['id'])
        fila = [caso, nid] + [round(float(r[k]), DECIMALES_FUERZA)
                              for k in ('fx', 'fy', 'fz', 'mx', 'my', 'mz')]
        fila.append('maestro' if nid in maestros else (SI if nid in en_diafragma else NO))
        fila += [cuenta(nid, 0)] if xy_iguales else [cuenta(nid, 0), cuenta(nid, 1)]
        fila.append(cuenta(nid, 2))
        filas.append(fila)
    return filas, xy_iguales


def _columnas_equilibrio():
    cols = []
    for eje in ('Fx', 'Fy', 'Fz'):
        cols += [Columna('Aplicada %s' % eje, 'kN', F_KN,
                         nota='suma de las cargas del caso (nodales + repartidas)'),
                 Columna('Reacción %s' % eje, 'kN', F_KN,
                         nota='suma de reacciones con la regla por grado de libertad '
                              '(calcular.equilibrio)'),
                 Columna('Error %s' % eje, 'kN', F_ERROR,
                         nota='aplicada + reacción; cero salvo el redondeo')]
    cols.append(Columna('Equilibrio confiable',
                        nota='no = alguna carga repartida no se pudo pasar a ejes globales'))
    return cols


def _filas_equilibrio(eq):
    fila = []
    for i in range(3):
        fila += [float(eq['aplicada_kN'][i]), float(eq['reaccion_kN'][i]),
                 float(eq['error_kN'][i])]
    fila.append(SI if eq.get('confiable', True) else
                'no (%d cargas sin convertir)' % int(eq.get('cargas_sin_convertir', 0)))
    return fila


def _nodo_de_maximo(desplazamientos):
    """(norma en mm, nodo) del mayor |u|; empate = el de menor id."""
    mejor = (-1.0, None)
    for d in sorted(desplazamientos, key=lambda d: int(d['id'])):
        n = _norma_mm(d)
        if n > mejor[0]:
            mejor = (n, int(d['id']))
    return mejor


# ------------------------------------------------------------
# Hojas de texto
# ------------------------------------------------------------
def _hoja_texto(nombre, filas, descripcion, columnas=None, color=None):
    columnas = columnas or [Columna('Tema', ancho=30), Columna('Qué dice', ancho=110)]
    return Hoja(nombre, columnas, filas, descripcion=descripcion, congelar_columnas=0,
                filtro=False, texto=True, color_pestana=color)


def _bloque(tema, lineas):
    """Filas [tema, linea]: el tema solo en la primera, para que se lea."""
    return [[tema if i == 0 else None, l] for i, l in enumerate(lineas)]


def _indice_de_hojas(hojas):
    filas = []
    for h in hojas:
        filas.append([Enlace(h.nombre, h.nombre),
                      '%s (%d filas de datos)' % (h.descripcion, len(h.filas))])
    return filas


def _columnas_con_nota(hojas):
    filas = []
    for h in hojas:
        notas = ['%s: %s' % (c.encabezado, c.nota) for c in h.columnas if c.nota]
        if notas:
            filas += _bloque('Columnas de %s' % h.nombre, notas)
    return filas


COMO_LEER = [
    'La fila 1 de cada hoja tiene el nombre de la columna y su unidad entre corchetes: '
    '"uz [mm]" son milimetros.',
    'La fila 1 queda fija al bajar. En las tablas tambien quedan fijas, al moverse a la '
    'derecha, las primeras columnas (Caso y el numero de nodo o elemento).',
    'Las flechas de la fila 1 son FILTROS: clic en la flecha de "Caso", deja marcado solo '
    'G, y la hoja muestra solo ese caso. Para ordenar, la misma flecha: "Ordenar de menor '
    'a mayor".',
    'Para volver a ver todo: pestaña Datos > Borrar (filtro).',
    'Los decimales se muestran con coma o con punto segun la configuracion de Windows; la '
    'celda guarda el numero con todos sus decimales (se ven al hacer clic en ella).',
    'En la hoja LEEME, el nombre de cada hoja es un enlace: clic para ir.',
]


# ------------------------------------------------------------
# El libro del anexo
# ------------------------------------------------------------
def _equilibrios_del_anexo(anexo, ctx):
    r"""
    {caso: (equilibrio, reacciones)} de cada caso y combinacion. Los
    casos base con lo que resolvio OpenSees; las combinaciones con
    combinar.combinar_cargas + combinar_resultados, la misma suma lineal
    que verifica comun/combinar.py. El equilibrio SIEMPRE lo calcula
    calcular.equilibrio: una sola regla.
    """
    calcular = _calcular()
    import combinar
    datos, resultados, casos = ctx['datos'], ctx['resultados'], ctx['arm']['casos']
    salida = {}
    for c in anexo['casos']:
        nombre = c['nombre']
        if c['tipo'] == 'caso' and nombre in resultados:
            caso, res = casos[nombre], resultados[nombre]
        else:
            activos = {k: float(v) for k, v in zip(CASOS_BASE, c['factores']) if v != 0.0}
            caso = combinar.combinar_cargas(datos, activos, nombre)
            res = combinar.combinar_resultados(resultados, activos)
        salida[nombre] = (calcular.equilibrio(datos, caso, res), res['reacciones'])
    return salida


def _resumen_demandas(caso):
    """Como VisorSemana04.ContarDemandas: NO PASA = !pasa (incluye las
    de P fuera de la curva). Peor D/C entre las que tienen D/C."""
    total = len(caso['demandas'])
    no_pasan = sum(1 for d in caso['demandas'] if not d['pasa'])
    fuera = sum(1 for d in caso['demandas'] if float(d['u']) >= U_FUERA_DE_CURVA)
    con_dc = [d for d in caso['demandas'] if float(d['u']) < U_FUERA_DE_CURVA]
    if con_dc:
        peor = max(sorted(con_dc, key=lambda d: int(d['id'])), key=lambda d: float(d['u']))
        return total, no_pasan, fuera, float(peor['u']), int(peor['id'])
    return total, no_pasan, fuera, None, None


def _hoja_resumen_anexo(anexo, equilibrios):
    columnas = ([Columna('Caso'),
                 Columna('Tipo', nota='caso base o combinación de casos base'),
                 Columna('Descripción', ancho=40)]
                + [Columna('λ%s' % k, '-', F_LAMBDA, nota='factor del caso %s' % k)
                   for k in CASOS_BASE]
                + _columnas_equilibrio()
                + [Columna('|u| máx', 'mm', F_MM,
                           nota='la mayor NORMA raiz(ux²+uy²+uz²) entre todos los nodos '
                                '(max_desplazamiento_mm del anexo de Unity)'),
                   Columna('Nodo del |u| máx', formato=F_ENTERO),
                   Columna('Corte basal Vx', 'kN', F_KN,
                           nota='|Reacción Fx|: el corte en la base en X'),
                   Columna('Corte basal Vy', 'kN', F_KN),
                   Columna('Barras revisadas D/C', formato=F_ENTERO,
                           nota='elementos con enfierradura (tienen curva P-M)'),
                   Columna('NO PASA', formato=F_ENTERO,
                           nota='D/C > 1, incluidas las de P fuera de la curva (lo mismo que '
                                'cuenta la cabecera del visor)'),
                   Columna('P fuera de la curva', formato=F_ENTERO),
                   Columna('Peor D/C', '-', F_DC,
                           nota='el mayor M/Mn entre las barras que tienen D/C'),
                   Columna('Elemento del peor D/C', formato=F_ENTERO)])
    filas = []
    for c in anexo['casos']:
        eq, _r = equilibrios[c['nombre']]
        total, no_pasan, fuera, peor, peor_id = _resumen_demandas(c)
        _n, nodo_max = _nodo_de_maximo(c['desplazamientos'])
        filas.append([c['nombre'], 'caso base' if c['tipo'] == 'caso' else 'combinación',
                      c['descripcion']]
                     + [float(v) for v in c['factores']]
                     + _filas_equilibrio(eq)
                     + [float(c['max_desplazamiento_mm']), nodo_max,
                        abs(float(eq['reaccion_kN'][0])), abs(float(eq['reaccion_kN'][1])),
                        total, no_pasan, fuera, peor, peor_id])
    reglas = [ReglaColor('Peor D/C', 'mayor', DC_ROJO, color=ROJO),
              ReglaColor('Peor D/C', 'entre', DC_AMBAR, DC_ROJO, color=AMBAR)]
    return Hoja('Resumen', columnas, filas, reglas=reglas,
                descripcion='Una fila por caso y combinación: factores, equilibrio por eje, '
                            'desplazamiento máximo, corte basal y demanda-capacidad')


def _hoja_elementos_anexo(anexo, modelo):
    area = {int(e['id']): float(e.get('area_tributaria') or 0.0) for e in modelo['elementos']}
    columnas = [
        Columna('Elemento', formato=F_ENTERO), Columna('Tipo'), Columna('Sección'),
        Columna('Nodo i', formato=F_ENTERO), Columna('Nodo j', formato=F_ENTERO),
        Columna('L', 'm', F_M), Columna('b', 'm', F_M), Columna('h', 'm', F_M),
        Columna('A', 'm²', F_AREA), Columna('Iy', 'm⁴', F_INERCIA),
        Columna('Iz', 'm⁴', F_INERCIA), Columna('J', 'm⁴', F_INERCIA),
        Columna('E', 'kPa', F_KPA), Columna('G', 'kPa', F_KPA),
        Columna("f'c", 'MPa', F_MPA, nota='vacio si la seccion trae E y G propios'),
        Columna('Material', ancho=40),
        Columna('vecxz', nota='vector que orienta los ejes locales (OpenSees geomTransf)'),
        Columna('Momento en el plano', nota='muros: el momento (My o Mz) del eje de '
                                            'inercia mayor, el que se compara con la P-M'),
        Columna('Área tributaria', 'm²', F_AREA, nota='losa que carga la barra; vacio = ninguna'),
        Columna('Cargada en G o Q'),
        Columna('Familia P-M', formato=F_ENTERO, nota='ver hoja Curvas P-M; vacio = sin fierro'),
        Columna('Condiciones de borde', ancho=50),
        Columna('Objeto en Unity', nota='nombre del objeto en el visor'),
        Columna('Línea de OpenSees', ancho=50),
    ]
    filas = []
    for e in anexo['elementos']:
        eid = int(e['id'])
        filas.append([
            eid, e['tipo'], e['seccion'], int(e['n1']), int(e['n2']),
            float(e['L']), float(e['b']) or None, float(e['h']) or None,
            float(e['A']), float(e['Iy']), float(e['Iz']), float(e['J']),
            float(e['E_kPa']), float(e['G_kPa']), float(e['fpc_MPa']) or None,
            e['material'], _texto_vector(e['vecxz']), e['momento_en_el_plano'] or None,
            area.get(eid) or None, SI if e['cargada'] else NO,
            int(e['familia']) if int(e['familia']) >= 0 else None,
            e['condiciones'], e['objeto_unity'], e['tag_opensees']])
    return Hoja('Elementos', columnas, filas,
                descripcion='Una barra por fila: nodos, sección, rigideces tal como las '
                            'recibe OpenSees, familia P-M y dónde está en Unity')


def _hoja_esfuerzos_anexo(anexo):
    tipo = {int(e['id']): e['tipo'] for e in anexo['elementos']}
    filas = []
    for c in anexo['casos']:
        for s in c['esfuerzos']:
            eid = int(s['id'])
            for extremo, k in (('i', 0), ('j', -1)):
                filas.append([c['nombre'], eid, tipo.get(eid), extremo, float(s['x'][k])]
                             + [float(s[m][k]) for m in ('N', 'Vy', 'Vz', 'T', 'My', 'Mz')])
    return Hoja('Esfuerzos', _columnas_esfuerzos(), filas, congelar_columnas=2,
                descripcion='Esfuerzos internos en los dos extremos de cada barra, por caso y '
                            'combinación (tracción positiva, convención del anexo)')


def _hoja_dc_anexo(anexo):
    por_id = {int(e['id']): e for e in anexo['elementos']}
    columnas = [
        Columna('Caso'), Columna('Elemento', formato=F_ENTERO), Columna('Tipo'),
        Columna('Sección'), Columna('Familia P-M', formato=F_ENTERO),
        Columna('Extremo', nota='el extremo con mayor momento'),
        Columna('P', 'kN', F_KN, nota='axial de demanda, COMPRESION positiva'),
        Columna('M', 'kN·m', F_KNM, nota='columna: raiz(My² + Mz²); muro: el momento de '
                                         'su plano'),
        Columna('M fuera del plano', 'kN·m', F_KNM, nota='solo muros; se informa, no se revisa'),
        Columna('Mn', 'kN·m', F_KNM, nota='capacidad NOMINAL de la curva P-M en esa P, sin '
                                          'factor φ. Vacio si P cae fuera de la curva'),
        Columna('D/C', '-', F_DC, nota='M / Mn. Vacio si P cae fuera de la curva. Rojo > 1, '
                                       'ámbar entre 0.9 y 1'),
        Columna('Estado', nota='pasa (D/C <= 1), NO PASA, o NO PASA: P fuera de la curva'),
    ]
    filas = []
    for c in anexo['casos']:
        for d in c['demandas']:
            eid = int(d['id'])
            e = por_id[eid]
            u = float(d['u'])
            fuera = u >= U_FUERA_DE_CURVA
            filas.append([
                c['nombre'], eid, e['tipo'], e['seccion'], int(d['familia']), d['extremo'],
                float(d['P']), float(d['M']),
                float(d['M_fuera_plano']) if e['tipo'] == 'muro' else None,
                None if fuera else float(d['Mn']),
                None if fuera else u,
                ESTADO_FUERA if fuera else (ESTADO_PASA if d['pasa'] else ESTADO_NO_PASA)])
    reglas = [ReglaColor('D/C', 'mayor', DC_ROJO, color=ROJO),
              ReglaColor('D/C', 'entre', DC_AMBAR, DC_ROJO, color=AMBAR)]
    return Hoja('Demanda-capacidad', columnas, filas, reglas=reglas, congelar_columnas=2,
                descripcion='Punto de demanda (P, M) de cada barra con fierro contra la '
                            'capacidad nominal Mn de su curva P-M, por caso y combinación')


def _hoja_curvas(anexo):
    columnas = [
        Columna('Familia P-M', formato=F_ENTERO), Columna('Punto', formato=F_ENTERO),
        Columna('Tipo'), Columna('Sección'),
        Columna('P', 'kN', F_KN, nota='compresion positiva'),
        Columna('Mn', 'kN·m', F_KNM, nota='momento nominal de la seccion de fibras en esa P'),
        Columna('Mmax', 'kN·m', F_KNM),
        Columna('Controla', nota='que limita ese punto de la curva'),
        Columna('As', 'cm²', F_CM2), Columna('Cuantía', '%', F_PCT),
        Columna('Refuerzo', ancho=45), Columna('Fuente del fierro', ancho=35),
        Columna('Barras de la familia', formato=F_ENTERO),
        Columna('Elementos', ancho=50),
        Columna('Clave', ancho=60, nota='lo que define la familia: dos barras con la misma '
                                        'clave comparten curva'),
    ]
    filas = []
    for fam in anexo['familias']:
        # Una sola barra va como numero: '28' escrito como texto sale en
        # Excel con el triangulo verde de "numero almacenado como texto".
        ids = (int(fam['elementos'][0]) if len(fam['elementos']) == 1
               else ', '.join(str(int(i)) for i in fam['elementos']))
        for k in range(len(fam['P'])):
            filas.append([int(fam['indice']), k + 1, fam['tipo'], fam['seccion'],
                          float(fam['P'][k]), float(fam['Mn'][k]), float(fam['Mmax'][k]),
                          fam['de'][k], float(fam['As_cm2']), float(fam['cuantia_pct']),
                          fam['refuerzo'], fam['fuente'], len(fam['elementos']), ids,
                          fam['clave']])
    return Hoja('Curvas P-M', columnas, filas,
                descripcion='Los puntos de cada curva de interacción P-M, una familia por '
                            'enfierradura distinta')


def _parrafos(lineas):
    """Las listas '_...' de parametros.json en parrafos: una linea vacia
    separa parrafos; las demas se juntan con espacio."""
    if isinstance(lineas, str):
        return lineas
    parrafos, actual = [], []
    for l in lineas or []:
        if str(l).strip():
            actual.append(str(l).strip())
        elif actual:
            parrafos.append(' '.join(actual))
            actual = []
    if actual:
        parrafos.append(' '.join(actual))
    return '\n'.join(parrafos)


def _hoja_supuestos_anexo(ed, anexo, ctx):
    import parametros
    p = ctx['p']
    crudo = {}
    if os.path.isfile(PARAMETROS_JSON):
        with io.open(PARAMETROS_JSON, encoding='utf-8') as f:
            crudo = json.load(f)
    cv, sismo = crudo.get('carga_viva', {}), crudo.get('sismo', {})
    filas = [
        ['Sobrecarga de uso q_Q', '%.4g kN/m² (%s)' % (p['q_Q'], parametros.origen_q(p)),
         _parrafos(cv.get('_q_kNm2'))],
        ['Coeficiente sísmico', 'V = %.4g · W' % p['coef_sismico'],
         _parrafos(sismo.get('_coeficiente'))],
        ['Peso sísmico', 'W = G + %.2f · Q' % p['fraccion_Q_sismica'],
         _parrafos(sismo.get('_fraccion_Q'))],
        ['Reparto del sismo en altura', parametros.texto_patron(p),
         _parrafos(sismo.get('_patron'))],
    ]
    for combo in p['combinaciones']:
        filas.append(['Combinación %s' % combo['nombre'], parametros.como_texto(combo),
                      _parrafos(combo.get('_por_que', ''))])
    if p['combinacion']['nombre'] not in [c['nombre'] for c in p['combinaciones']]:
        filas.append(['Combinación %s' % p['combinacion']['nombre'],
                      parametros.como_texto(p['combinacion']), 'dada por línea de comandos'])
    filas.append(['Combinaciones revisadas',
                   'solo las de arriba (%s)' % ', '.join(c['nombre'] for c in p['combinaciones']),
                   'Son las declaradas en semana03/parametros.json. No es una envolvente de '
                   'diseño: una combinación que no está no se revisa.'])

    modelo = ctx['modelo']
    origenes, aceros = {}, {}
    for e in modelo['elementos']:
        fe = e.get('enfierradura') or {}
        o = (fe.get('longitudinal') or {}).get('origen')
        if o:
            origenes[o] = origenes.get(o, 0) + 1
        a = (fe.get('acero') or {}).get('designacion')
        if a:
            aceros[a] = aceros.get(a, 0) + 1
    for o in sorted(origenes):
        filas.append(['Fierro longitudinal', '%d elementos: %s' % (origenes[o], o),
                      'Campo "origen" de la enfierradura en data/modelo/%s.json' % ed])
    for a in sorted(aceros):
        filas.append(['Acero de refuerzo', '%s en %d elementos' % (a, aceros[a]),
                      'Campo "acero" de la enfierradura en data/modelo/%s.json' % ed])
    mat = modelo.get('material', {})
    propias = sum(1 for s in modelo['secciones'] if s.get('E'))
    filas.append(['Hormigón', "f'c = %g MPa del material del modelo; %d secciones traen E y G "
                              "propios" % (float(mat.get('fpc_MPa', 0.0)), propias),
                  'Ec = 4700·raiz(f\'c)·1000 kPa (comun/servidor_opensees.py). Las secciones '
                  'con E propio son las de un cuerpo con otro hormigón (conjunto).'])
    filas += [
        ['Capacidad', 'nominal, sin factor φ',
         'Mn sale de la sección de fibras de comun/capacidad.py. D/C = M/Mn compara, no '
         'diseña.'],
        ['Redondeo', 'desplazamientos a 8 decimales (m), fuerzas a 4 (kN)',
         'Lo hace el servidor (servidor_opensees.extraer_resultados). Diferencias por '
         'debajo de eso no significan nada.'],
        ['Fuente de los resultados', anexo['elementos'][0]['resultados'] if anexo['elementos']
         else '', 'Los data/resultados/<ed>_<caso>.json usan la Q y el sismo del modelo, no '
                  'estos parámetros: sus números no tienen por qué coincidir con este libro.'],
        ['Desplazamiento máximo', '|u| máx = la mayor NORMA de la traslación',
         'El servidor de reanálisis (/analizar) informa la mayor COMPONENTE (|ux|, |uy| o '
         '|uz|): son dos números distintos a propósito.'],
    ]
    cuerpos = (modelo.get('info') or {}).get('cuerpos') or [ed]
    perfiles = []
    for cuerpo in cuerpos:
        carpeta = os.path.join(rutas.edificio(cuerpo), 'perfiles')
        if os.path.isdir(carpeta):
            perfiles += ['edificios/%s/perfiles/%s' % (cuerpo, n)
                         for n in sorted(os.listdir(carpeta)) if n.endswith('.json')]
    if perfiles:
        filas.append(['Lo que el plano no dice', ', '.join(perfiles),
                      'Cada supuesto del modelo (dinteles, diámetros, junta, cotas) está en '
                      'esos archivos con su justificación al lado.'])
    columnas = [Columna('Tema', ancho=28), Columna('Supuesto', ancho=45),
                Columna('Por qué / de dónde sale', ancho=90)]
    return Hoja('Supuestos', columnas, filas, congelar_columnas=0, filtro=False, texto=True,
                descripcion='Lo que no sale del plano ni de OpenSees: parámetros del '
                            'profesor, combinaciones, fierro supuesto, capacidad nominal')


def _comando(ed, argv):
    return ' '.join([COMANDO_CLI, ed] + [str(a) for a in (argv or [])])


def _hoja_leeme_anexo(ed, anexo, ctx, argv, hojas):
    info, modelo = anexo['info'], ctx['modelo']
    n_base = sum(1 for c in anexo['casos'] if c['tipo'] == 'caso')
    filas = []
    filas += _bloque('Qué es este archivo', [
        'Los resultados del edificio %s calculados con OpenSees desde Python: los mismos '
        'números que muestra el visor de Unity (anexo semana04.json) cuando se exporta con '
        'los mismos parámetros.' % ed.upper(),
        'Ninguna celda tiene fórmulas: todo lo escribió Python. Si cambias un valor, no se '
        'recalcula nada. Para otros parámetros se vuelve a generar (abajo).',
    ])
    filas += _bloque('Cómo se regenera', [
        'Desde la carpeta del repositorio, en una terminal:',
        _comando(ed, argv),
        'Con otros parámetros se agregan al final, p. ej. --cs 0.20 o --q 2.5 (los de '
        'semana03/parametros.py).',
        'Cierra este archivo en Excel antes: Windows no deja reemplazar un archivo abierto.',
        r'Se comprueba con: .venv\Scripts\python.exe semana05\test_excel.py %s' % ed,
    ])
    filas += _bloque('Edificio', [
        '%s: %s' % (ed, (modelo.get('info') or {}).get('descripcion', '')),
        '%d nodos, %d elementos, %d pisos rígidos (diafragmas)'
        % (len(modelo['nodos']), len(modelo['elementos']), len(modelo.get('diafragmas', []))),
        '%d casos base (%s) y %d combinaciones (%s)'
        % (n_base, ', '.join(c['nombre'] for c in anexo['casos'] if c['tipo'] == 'caso'),
           len(anexo['casos']) - n_base,
           ', '.join(c['nombre'] for c in anexo['casos'] if c['tipo'] != 'caso')),
        '%d familias P-M para las barras con enfierradura' % len(anexo['familias']),
    ])
    filas += _bloque('Parámetros', info['parametros'])
    filas += _bloque('Huella de la fuente', [
        'data/modelo/%s.json  sha256 %s' % (ed, _huella(rutas.modelo(ed))),
        'semana03/parametros.json  sha256 %s' % _huella(PARAMETROS_JSON),
        'Si alguno cambió, este libro está desactualizado: regenerarlo.',
    ])
    filas += _bloque('Cómo leer una hoja', COMO_LEER)
    filas += _bloque('Unidades', [
        'Longitudes y coordenadas en m; desplazamientos en mm (en los JSON están en m); '
        'giros en rad.',
        'Fuerzas en kN; momentos en kN·m; módulos E y G en kPa; f\'c en MPa.',
    ])
    filas += _bloque('Ejes y signos', [
        'Ejes de OpenSees: Z vertical hacia arriba (Unity usa Y vertical).',
        'Desplazamientos: positivos en el sentido de los ejes globales.',
        'Reacciones: la fuerza que el apoyo hace sobre la estructura. Por eso aplicada + '
        'reacción = 0 (columna Error del Resumen).',
        'Esfuerzos: internos, en ejes locales de la barra, TRACCIÓN positiva. Convención del '
        'anexo:',
    ] + ['    ' + l for l in info['convencion']])
    filas += _bloque('Advertencias', [
        '1. NO sumes la columna entera de Reacciones. Los nodos de un piso rígido traen la '
        'fuerza interna del diafragma y el corte basal saldría al doble o más. Filtra "Cuenta '
        'en Fx/Fy" = sí para Fx y Fy, o "Cuenta en Fz" = sí para Fz. Los totales correctos ya '
        'están en la hoja Resumen.',
        '2. D/C es NOMINAL: M / Mn con Mn sin factor φ. Sirve para comparar barras, no es una '
        'verificación de diseño.',
        '3. D/C vacío con Estado "%s": la carga axial cae fuera de la curva P-M de su familia '
        '(ahí Mn = 0). Cuenta como NO PASA.' % ESTADO_FUERA,
        '4. Colores de D/C: rojo = mayor que 1, ámbar = entre 0.9 y 1.',
        '5. |u| máx es la NORMA raiz(ux²+uy²+uz²). El servidor de reanálisis informa la mayor '
        'componente: son números distintos.',
        '6. Solo están las combinaciones declaradas (hoja Supuestos): no es una envolvente de '
        'diseño.',
    ])
    filas += _bloque('Hojas', [])
    filas += _indice_de_hojas(hojas)
    filas += _columnas_con_nota(hojas)
    return _hoja_texto('LEEME', filas, 'Esta hoja', color='70AD47')


def hojas_del_anexo(ed, anexo, ctx, argv=None):
    r"""
    Las hojas del libro del edificio, sin escribir nada. Separado de
    escribir_libro_anexo para que el test pueda armar el anexo UNA vez
    y comparar contra el mismo.
    """
    equilibrios = _equilibrios_del_anexo(anexo, ctx)
    modelo = ctx['modelo']
    cota = {int(n['id']): float(n['z']) for n in modelo['nodos']}

    desplazamientos = []
    for c in anexo['casos']:
        desplazamientos += _filas_desplazamientos(c['nombre'], c['desplazamientos'], cota)

    reacciones, xy_iguales = [], True
    for c in anexo['casos']:
        filas, xy_iguales = _filas_reacciones(c['nombre'], equilibrios[c['nombre']][1], modelo)
        reacciones += filas

    hojas = [
        _hoja_resumen_anexo(anexo, equilibrios),
        _hoja_nodos(modelo),
        Hoja('Desplazamientos', _columnas_desplazamientos(), desplazamientos,
             congelar_columnas=2,
             descripcion='Desplazamientos y giros de cada nodo, por caso y combinación'),
        _hoja_elementos_anexo(anexo, modelo),
        _hoja_esfuerzos_anexo(anexo),
        Hoja('Reacciones', _columnas_reacciones(xy_iguales), reacciones, congelar_columnas=2,
             descripcion='Reacciones de cada nodo restringido, por caso y combinación. NO '
                         'sumar la columna entera: filtrar por "Cuenta en ..."'),
        _hoja_dc_anexo(anexo),
        _hoja_curvas(anexo),
        _hoja_supuestos_anexo(ed, anexo, ctx),
    ]
    return [_hoja_leeme_anexo(ed, anexo, ctx, argv, hojas)] + hojas


def escribir_libro_desde_anexo(ed, anexo, ctx, ruta, argv=None):
    """El libro del edificio desde un anexo ya construido. Devuelve la
    ruta absoluta."""
    hojas = hojas_del_anexo(ed, anexo, ctx, argv)
    return escribir(ruta, hojas,
                    titulo='Resultados %s - laboratorio estructural Grupo 7' % ed.upper(),
                    asunto='Casos, combinaciones, equilibrio, esfuerzos y demanda-capacidad',
                    descripcion=_comando(ed, argv))


def escribir_libro_anexo(ed, ruta, argv=None) -> str:
    r"""
    Escribe el libro de resultados de un edificio y devuelve la ruta
    ABSOLUTA del .xlsx escrito (str).

    ed    'lt2', 'ingenieria' o 'conjunto' (otro: ValueError).
    ruta  el .xlsx a escribir; crea la carpeta si falta. El CLI
          (semana05/exportar_excel.py) pasa rutas.excel_resultados(ed).
    argv  lista de flags de semana03/parametros.py (['--cs', '0.20']).
          None = los parametros por defecto.

    Fuente: construir_anexo(ed, argv) de semana04/exportar_unity.py,
    cargado POR RUTA, mas calcular.equilibrio por caso y combinacion.
    Nunca lee data/unity/semana04.json (guarda solo el ultimo edificio
    exportado).

    Errores: PermissionError con mensaje claro si el archivo esta
    abierto en Excel (se comprueba ANTES de calcular); ValueError si el
    edificio no existe.
    """
    if ed not in EDIFICIOS:
        raise ValueError('edificio %r: use %s' % (ed, ', '.join(EDIFICIOS)))
    argv = list(argv or [])
    comprobar_que_se_puede_escribir(ruta)
    eu = _exportador_semana04()
    anexo, ctx = eu.construir_anexo(ed, argv)
    return escribir_libro_desde_anexo(ed, anexo, ctx, ruta, argv)


# ------------------------------------------------------------
# El libro del reanalisis
# ------------------------------------------------------------
def _casos_de_la_respuesta(modelo, respuesta):
    r"""
    [(caso de carga, resultado)] en las dos formas del servidor. La
    multi-caso trae 'casos' en el orden de modelo['casos_de_carga']; la
    plana (un caso) trae las listas en la raiz y las cargas en la raiz
    del modelo.
    """
    if respuesta.get('casos') is not None:
        cargas = list(modelo.get('casos_de_carga') or [])
        por_nombre = {c.get('nombre'): c for c in cargas}
        salida = []
        for i, r in enumerate(respuesta['casos']):
            caso = por_nombre.get(r.get('nombre'))
            if caso is None and i < len(cargas):
                caso = cargas[i]
            salida.append((caso or {'nombre': r.get('nombre', 'caso_%d' % (i + 1))}, r))
        return salida
    if respuesta.get('desplazamientos') is None:
        return []
    caso = {'nombre': modelo.get('nombre_caso', 'unico'),
            'cargas_nodales': modelo.get('cargas_nodales', []),
            'cargas_distribuidas': modelo.get('cargas_distribuidas', [])}
    r = dict(respuesta)
    r.setdefault('nombre', caso['nombre'])
    return [(caso, r)]


def _equilibrio_de(modelo, caso, r):
    """(equilibrio, origen): el que trae la respuesta si vino completo; si
    no, calcular.equilibrio sobre el mismo modelo."""
    eq = r.get('equilibrio')
    if isinstance(eq, dict) and len(eq.get('aplicada_kN') or []) == 3:
        return eq, 'servidor'
    return _calcular().equilibrio(modelo, caso, r), 'calculado al escribir el libro'


def _largo(nodos, e):
    a, b = nodos.get(int(e['n1'])), nodos.get(int(e['n2']))
    if a is None or b is None:
        return None
    return math.sqrt(sum((float(b[k]) - float(a[k])) ** 2 for k in ('x', 'y', 'z')))


def hojas_del_reanalisis(modelo, respuesta):
    """Las hojas del libro de un reanalisis, sin escribir nada."""
    info = modelo.get('info') or {}
    edificio = (info.get('edificio') or '').strip() or '(sin nombre en info.edificio)'
    nodos = {int(n['id']): n for n in modelo['nodos']}
    cota = {nid: float(n['z']) for nid, n in nodos.items()}
    por_id = {int(e['id']): e for e in modelo['elementos']}
    casos = _casos_de_la_respuesta(modelo, respuesta)

    resumen, desplazamientos, esfuerzos, reacciones = [], [], [], []
    xy_iguales = True
    for caso, r in casos:
        nombre = r.get('nombre') or caso.get('nombre', '')
        desp = r.get('desplazamientos') or []
        eq, origen = _equilibrio_de(modelo, caso, r)
        maximo, nodo_max = _nodo_de_maximo(desp) if desp else (None, None)
        resumen.append([nombre, SI if r.get('ok', respuesta.get('ok', True)) else NO]
                       + _filas_equilibrio(eq)
                       + [origen,
                          _mm(r['max_desplazamiento']) if r.get('max_desplazamiento') is not None
                          else None,
                          maximo, nodo_max,
                          abs(float(eq['reaccion_kN'][0])), abs(float(eq['reaccion_kN'][1]))])
        desplazamientos += _filas_desplazamientos(nombre, desp, cota)
        filas, xy_iguales = _filas_reacciones(nombre, r.get('reacciones') or [], modelo)
        reacciones += filas
        for fe in sorted(r.get('fuerzas_elementos') or [], key=lambda x: int(x['id'])):
            eid = int(fe['id'])
            f = [float(v) for v in fe['f']]
            e = por_id.get(eid)
            tipo = e.get('tipo', '') if e else None
            L = _largo(nodos, e) if e else None
            # Esfuerzo interno con traccion positiva: -f en i, +f en j
            # (los mismos de exportar_unity.esfuerzos_internos en x = 0 y x = L).
            esfuerzos.append([nombre, eid, tipo, 'i', 0.0] + [-v for v in f[:6]])
            esfuerzos.append([nombre, eid, tipo, 'j', L] + f[6:12])

    col_resumen = ([Columna('Caso'), Columna('Resolvió', nota='no = OpenSees no convergió')]
                   + _columnas_equilibrio()
                   + [Columna('Origen del equilibrio',
                              nota='servidor = lo trajo /analizar; calculado = '
                                   'calcular.equilibrio al escribir el libro'),
                      Columna('Mayor componente', 'mm', F_MM,
                              nota='max_desplazamiento del servidor: el mayor |ux|, |uy| o |uz|'),
                      Columna('|u| máx', 'mm', F_MM, nota='la mayor NORMA raiz(ux²+uy²+uz²)'),
                      Columna('Nodo del |u| máx', formato=F_ENTERO),
                      Columna('Corte basal Vx', 'kN', F_KN), Columna('Corte basal Vy', 'kN', F_KN)])

    secciones = modelo.get('secciones') or []
    if isinstance(secciones, dict):
        # El servidor acepta tambien {nombre: seccion} (normalizar_secciones);
        # el contrato dice lista, pero un modelo resuelto no se queda sin
        # libro por la forma.
        secciones = [dict(s, nombre=n) for n, s in secciones.items()]
    secciones = {s.get('nombre'): s for s in secciones}
    mat = modelo.get('material') or {}
    elementos = []
    for e in sorted(modelo['elementos'], key=lambda e: int(e['id'])):
        s = secciones.get(e.get('seccion'), {})
        # En lo que manda Unity, E = 0 o G = 0 es "no vino" (JsonUtility
        # escribe todos los campos): se usa el material del modelo.
        propio = float(s.get('E') or 0.0) > 0.0
        elementos.append([
            int(e['id']), e.get('tipo', ''), e.get('seccion', ''), int(e['n1']), int(e['n2']),
            _largo(nodos, e), float(s['A']) if 'A' in s else None,
            float(s['Iy']) if 'Iy' in s else None, float(s['Iz']) if 'Iz' in s else None,
            float(s['J']) if 'J' in s else None,
            float(s['E']) if propio else None,
            float(s['G']) if float(s.get('G') or 0.0) > 0.0 else None,
            'E y G propios de la sección' if propio else
            "material del modelo: f'c %g MPa" % float(mat.get('fpc_MPa', 0.0))])
    col_elementos = [
        Columna('Elemento', formato=F_ENTERO), Columna('Tipo'), Columna('Sección'),
        Columna('Nodo i', formato=F_ENTERO), Columna('Nodo j', formato=F_ENTERO),
        Columna('L', 'm', F_M), Columna('A', 'm²', F_AREA), Columna('Iy', 'm⁴', F_INERCIA),
        Columna('Iz', 'm⁴', F_INERCIA), Columna('J', 'm⁴', F_INERCIA),
        Columna('E', 'kPa', F_KPA, nota='vacio = el del material del modelo'),
        Columna('G', 'kPa', F_KPA), Columna('Material', ancho=35)]

    hojas = [
        Hoja('Resumen', col_resumen, resumen,
             descripcion='Una fila por caso: equilibrio por eje, desplazamiento máximo y corte '
                         'basal'),
        _hoja_nodos(modelo),
        Hoja('Desplazamientos', _columnas_desplazamientos(), desplazamientos,
             congelar_columnas=2, descripcion='Desplazamientos y giros de cada nodo, por caso'),
        Hoja('Elementos', col_elementos, elementos,
             descripcion='Las barras del modelo tal como llegaron al servidor'),
        Hoja('Esfuerzos', _columnas_esfuerzos(), esfuerzos, congelar_columnas=2,
             descripcion='Esfuerzos internos en los extremos de cada barra (tracción positiva)'),
        Hoja('Reacciones', _columnas_reacciones(xy_iguales), reacciones, congelar_columnas=2,
             descripcion='Reacciones de cada nodo restringido. NO sumar la columna entera: '
                         'filtrar por "Cuenta en ..."'),
    ]

    avisos = [str(a) for a in (respuesta.get('avisos') or [])]
    filas = []
    filas += _bloque('Qué es este archivo', [
        'El resultado del ÚLTIMO reanálisis pedido desde Unity al servidor (POST /analizar): '
        'el modelo tal como estaba en el visor, con las ediciones hechas ahí.',
        'Se sobrescribe en cada reanálisis. No es el libro versionado del edificio '
        '(data/excel/<edificio>_resultados.xlsx), que usa los parámetros del laboratorio.',
        'Ninguna celda tiene fórmulas: todo lo escribió Python.',
    ])
    filas += _bloque('Modelo', [
        'edificio: %s' % edificio,
        '%d nodos, %d elementos, %d pisos rígidos; %d casos resueltos (%s)'
        % (len(modelo['nodos']), len(modelo['elementos']), len(modelo.get('diafragmas') or []),
           len(casos), ', '.join(str(r.get('nombre', '')) for _c, r in casos)),
        'el servidor respondió ok = %s' % (SI if respuesta.get('ok', True) else NO),
    ])
    if avisos:
        filas += _bloque('Avisos del servidor', avisos)
    filas += _bloque('Cómo leer una hoja', COMO_LEER)
    filas += _bloque('Unidades y signos', [
        'Coordenadas en m; desplazamientos en mm (el servidor los da en m); giros en rad; '
        'fuerzas en kN; momentos en kN·m.',
        'Ejes de OpenSees: Z vertical. Reacciones = fuerza del apoyo sobre la estructura '
        '(aplicada + reacción = 0).',
        'Esfuerzos: internos, en ejes locales, TRACCIÓN positiva: en i es -f_i y en j es +f_j '
        'de eleResponse(localForce).',
    ])
    filas += _bloque('Advertencias', [
        '1. NO sumes la columna entera de Reacciones: filtra "Cuenta en Fx/Fy" o "Cuenta en '
        'Fz" = sí. Los totales correctos están en Resumen.',
        '2. "Mayor componente" es lo que informa el servidor (el mayor |ux|, |uy| o |uz|); '
        '"|u| máx" es la norma. Son números distintos a propósito.',
    ])
    filas += _bloque('Hojas', [])
    filas += _indice_de_hojas(hojas)
    filas += _columnas_con_nota(hojas)
    return [_hoja_texto('LEEME', filas, 'Esta hoja', color='70AD47')] + hojas


def escribir_libro_reanalisis(modelo: dict, respuesta: dict, ruta) -> str:
    r"""
    Escribe el libro de UN reanalisis del servidor y devuelve la ruta
    ABSOLUTA del .xlsx escrito (str).

    modelo     dict: el JSON que Unity manda a POST /analizar, tal cual
               (JsonUtility.ToJson de ModeloEstructural).
    respuesta  dict: lo que devuelve servidor_opensees.construir_y_resolver,
               MULTI-CASO {ok, error, avisos, casos[{nombre, ok,
               max_desplazamiento, desplazamientos, reacciones,
               fuerzas_elementos, equilibrio}]} o PLANA (un caso, las
               listas en la raiz). Si un caso no trae 'equilibrio' (o
               viene null), se calcula con calcular.equilibrio.
    ruta       el .xlsx a escribir; el servidor pasa
               rutas.excel_reanalisis(<edificio>).

    Errores: los deja subir; el servidor los atrapa y los devuelve en
    'excel_error' sin romper /analizar.
    """
    hojas = hojas_del_reanalisis(modelo, respuesta)
    info = modelo.get('info') or {}
    return escribir(ruta, hojas,
                    titulo='Reanálisis %s - laboratorio estructural Grupo 7'
                           % ((info.get('edificio') or '').strip() or 'modelo'),
                    asunto='Resultado de POST /analizar',
                    descripcion='comun/servidor_opensees.py')
